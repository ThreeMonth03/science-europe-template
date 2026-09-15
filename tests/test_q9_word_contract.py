import copy
from pathlib import Path
import sys
import unittest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from q9_word_contract import expected_blocks,xml,text,separator


def document():
    d=Document()
    for name in ['Compact','Pilot List Lead']:d.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    d.add_paragraph('1. Original',style='Heading 3');d.add_paragraph('Author text.')
    d.add_paragraph('9. Ethics?',style='Heading 3');d.add_paragraph('Name.csv',style='Compact')
    d.add_paragraph('Personal.',style='Pilot List Lead');d.add_paragraph('Sensitive.',style='Compact')
    d.add_paragraph('10. Sharing?',style='Heading 3');return d


class Q9WordContractTests(unittest.TestCase):
    def test_sentence_spacing_is_language_appropriate(self):
        self.assertEqual('',separator('不包含個人資料。','不包含敏感資料。'))
        self.assertEqual(' ',separator('No personal data.','No sensitive data.'))

    def test_exact_projection_keeps_both_facts_and_other_blocks(self):
        old=document();before=[xml(n) for n in old.element.body]
        result,counts=expected_blocks(list(old.element.body),[('Name.csv',['Personal.','Sensitive.'])])
        self.assertEqual({'styled_labels':1,'joined_flag_pairs':1},counts)
        self.assertEqual('Personal. Sensitive.',text(result[4]))
        self.assertEqual(before[:3],[xml(n) for n in result[:3]])
        self.assertEqual(before[6:],[xml(n) for n in result[5:]])
        self.assertEqual(before,[xml(n) for n in old.element.body])

    def test_wrong_fact_order_or_missing_group_is_rejected(self):
        for plan in [[('Name.csv',['Sensitive.','Personal.'])],[('Missing.csv',['Personal.'])]]:
            with self.assertRaises(AssertionError):expected_blocks(list(document().element.body),plan)

    def test_formatted_authored_flags_cannot_be_flattened(self):
        d=document();d.paragraphs[4].runs[0].bold=True
        with self.assertRaises(AssertionError):expected_blocks(list(d.element.body),[('Name.csv',['Personal.','Sensitive.'])])
